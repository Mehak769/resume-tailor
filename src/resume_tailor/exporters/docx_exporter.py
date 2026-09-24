from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor

from resume_tailor.core.exceptions import ResumeExportError
from resume_tailor.core.models import TailoredResume
from resume_tailor.exporters.base import BaseResumeExporter


class DocxResumeExporter(BaseResumeExporter):
    def export(self, tailored_resume: TailoredResume, output_path: Path) -> Path:
        try:
            resolved_output = output_path.resolve()
            resolved_output.parent.mkdir(parents=True, exist_ok=True)
            doc = docx.Document()

            # Standard 0.75-inch margins
            for section in doc.sections:
                section.top_margin = Inches(0.75)
                section.bottom_margin = Inches(0.75)
                section.left_margin = Inches(0.75)
                section.right_margin = Inches(0.75)

            p_res = tailored_resume.parsed_resume
            contact = p_res.contact_info

            # Header / Name
            title_p = doc.add_paragraph()
            title_p.paragraph_format.space_after = Pt(2)
            name_run = title_p.add_run(contact.name or "Candidate Name")
            name_run.font.size = Pt(20)
            name_run.font.bold = True

            # Contact line
            contact_parts = [
                x
                for x in [
                    contact.email,
                    contact.phone,
                    contact.location,
                    contact.linkedin,
                    contact.github,
                ]
                if x
            ]
            if contact_parts:
                c_p = doc.add_paragraph(" | ".join(contact_parts))
                c_p.paragraph_format.space_after = Pt(10)

            # Professional Summary
            if p_res.summary:
                self._add_heading(doc, "Professional Summary")
                doc.add_paragraph(p_res.summary)

            # Skills
            if p_res.skills:
                self._add_heading(doc, "Skills & Competencies")
                doc.add_paragraph(", ".join(p_res.skills))

            # Experience
            if p_res.experience:
                self._add_heading(doc, "Professional Experience")
                for exp in p_res.experience:
                    exp_p = doc.add_paragraph()
                    exp_p.paragraph_format.space_before = Pt(6)
                    exp_p.paragraph_format.space_after = Pt(2)

                    role_run = exp_p.add_run(f"{exp.role}")
                    role_run.bold = True
                    exp_p.add_run(f" | {exp.company}")
                    if exp.start_date or exp.end_date:
                        date_str = f" ({exp.start_date} - {exp.end_date})"
                        exp_p.add_run(date_str)

                    for bullet in exp.bullet_points:
                        doc.add_paragraph(bullet, style="List Bullet")

            # Education
            if p_res.education:
                self._add_heading(doc, "Education")
                for edu in p_res.education:
                    edu_p = doc.add_paragraph()
                    edu_run = edu_p.add_run(f"{edu.degree}")
                    edu_run.bold = True
                    edu_p.add_run(f", {edu.institution}")
                    if edu.graduation_date:
                        edu_p.add_run(f" ({edu.graduation_date})")

            # Projects
            if p_res.projects:
                self._add_heading(doc, "Projects")
                for proj in p_res.projects:
                    proj_p = doc.add_paragraph()
                    proj_run = proj_p.add_run(proj.title)
                    proj_run.bold = True
                    if proj.technologies:
                        proj_p.add_run(f" [{', '.join(proj.technologies)}]")
                    for desc in proj.description:
                        doc.add_paragraph(desc, style="List Bullet")

            doc.save(str(resolved_output))
            return resolved_output

        except Exception as e:
            if isinstance(e, ResumeExportError):
                raise
            raise ResumeExportError(f"Failed to generate DOCX resume: {e}") from e

    def _add_heading(self, doc: DocxDocument, text: str) -> None:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2B, 0x2D, 0x42)
